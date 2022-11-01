import openpyxl
from openpyxl.chart import Reference, BarChart, LineChart, Series
from openpyxl.chart.label import DataLabelList

file = r'C:\Users\Артем\Desktop\Industry consumption.xlsx'

df = openpyxl.load_workbook(file)  # Читаем файл
sheet = df['Total']  # Выбираем нужный лист

chart = LineChart()  # Создаем объект LineChart

#countries = Reference(sheet, min_col=34, max_col=34, min_row=2, max_row=159)
years = Reference(sheet, min_col=35, max_col=64, min_row=1, max_row=1)  # Подаем список годов, по которому будет определяться ось х на графике
#data = Reference(sheet, min_col=35, max_col=64, min_row=2, max_row=159)
# Записываем легенду графика, а также определяем данные, по которым строится сам график
for i in range(2, 160):
    chart.series.append(
        Series(Reference(sheet, min_col=35, max_col=64, min_row=i, max_row=i), title=sheet.cell(i, 34).value))
# chart.add_data(data, from_rows=True)
chart.set_categories(years)  # Указываем, какой должна быть ось х на графике
chart.width = 30  # Ширина и высота графика (в см)
chart.height = 10

sheet.add_chart(chart, "E5")  # Добавляем график на лист в переменной sheet с левым верхним углом в ячейке Е5
df.save(file)